"""Read-only reconciliation. Output candidates/states, never authorize deletion by itself."""
import argparse,csv,hashlib,json,re
from collections import defaultdict
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from delivery_paths import pending_files, history_rows

def sha(p):
    with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()

def reconcile(settings,mapping,history_dir):
    pending=Path(settings['pending_root']);ledgerroot=Path(settings['ledger_root'])
    files=pending_files(settings)
    by_name=defaultdict(set);by_hash=defaultdict(set)
    def add(row):
        no=row.get('delivery_no');src=row.get('source_file');h=row.get('source_sha256')
        if not no or not src:return
        by_name[Path(src).name].add(no)
        if h and re.fullmatch('[a-fA-F0-9]{64}',h):by_hash[h.lower()].add(no)
    for row in mapping:add(row)
    for _, row in history_rows(history_dir):add(row)
    archives={'primary':defaultdict(list),'settlement':defaultdict(list)}
    for kind,key in [('primary','primary_archive_root'),('settlement','settlement_archive_root')]:
        for p in Path(settings[key]).rglob('*'):
            if p.is_file() and p.suffix.lower() in ('.jpg','.jpeg','.png'):
                no=re.sub(r'^\d{4}已录','',p.stem) if kind=='settlement' else p.stem
                archives[kind][no].append(p)
    item_ids={};source_hashes={}
    for p in files:
        h=sha(p);source_hashes[p.name]=h;item_ids[p.name]=by_name[p.name]|by_hash[h]
    ids=set().union(*item_ids.values()) if files else set()
    records=defaultdict(list);related=defaultdict(list);ledgers=[]
    for p in ledgerroot.glob('*.xlsx'):
        if p.name.startswith('~$') or '模板' in p.name:continue
        before=sha(p);wb=load_workbook(p,read_only=True,data_only=False)
        try:
            for ws in wb.worksheets:
                headers=None;col=None
                for rn,row in enumerate(ws.iter_rows(values_only=True),1):
                    if col is None:
                        if '送货单编号' not in row:continue
                        headers=[str(v or '').strip() for v in row];col=row.index('送货单编号');continue
                    no=row[col] if len(row)>col else None
                    suffix_bases=[base for base in ids if isinstance(no,str) and re.fullmatch(re.escape(base)+r'-[1-9]\d*',no)]
                    if no in ids or suffix_bases:
                        vals={headers[j]:v for j,v in enumerate(row) if j<len(headers) and headers[j]}
                        record={'ledger_file':str(p),'sheet':ws.title,'row':rn,'values':vals}
                        if no in ids:records[no].append(record)
                        for base in suffix_bases:related[base].append(record)
        finally:wb.close()
        after=sha(p)
        if before!=after:raise RuntimeError(f'Ledger changed during read-only scan: {p}')
        ledgers.append({'file':str(p),'sha256':after,'mtime':datetime.fromtimestamp(p.stat().st_mtime).astimezone().isoformat(),'unchanged_during_scan':True})
    out=[]
    for p in files:
        candidates=sorted(item_ids[p.name]);row={'source_file':str(p),'source_sha256':source_hashes[p.name],'candidate_numbers':candidates,'status':'identity_review','reason':'No unique reviewed number','ledger_records':[],'primary':[],'settlement':[]}
        row['related_suffix_records']=[r for no in candidates for r in related[no]]
        if row['related_suffix_records']:row['suffix_review_warning']='Compare all physical-note fields before choosing an existing suffix; never create or merge suffixes automatically.'
        if len(candidates)==1:
            no=candidates[0];rs=records[no];row['ledger_records']=rs
            for kind in archives:
                row[kind]=[{'path':str(q),'sha256':sha(q),'same_source':sha(q)==row['source_sha256']} for q in archives[kind][no]]
            if not rs:row.update(status='not_entered_or_number_conflict',reason='Exact full number absent; do not strip zeros/suffixes')
            elif len({r['ledger_file'] for r in rs})>1:row.update(status='cross_carrier_conflict',reason='Number appears in more than one ledger')
            else:
                missing=[]
                for r in rs:
                    for field in ['计划单日期','到货日期','采购货物名称','品牌','型号','规格','数量','件数']:
                        if r['values'].get(field) in (None,''):missing.append(f"{r['sheet']}!{r['row']}:{field}")
                row['missing_fields']=missing
                if missing:row.update(status='partial_entry',reason=';'.join(missing))
                elif any(x['same_source'] for x in row['primary']) and any(x['same_source'] for x in row['settlement']):row.update(status='dual_archive_pending_review_cleanup',reason='Both copies have source hash; still validate canonical target and ledger fields before cleanup')
                elif row['primary'] and row['settlement']:row.update(status='version_or_business_conflict',reason='Exact number exists but photo hashes differ; visual identity/field review required')
                elif row['primary'] or row['settlement']:row.update(status='single_archive',reason='One archive missing or differs; photo-to-ledger review required')
                else:row.update(status='saved_unarchived',reason='Rows exist; full photo-to-ledger review required before archive')
        out.append(row)
    return {'generated_at':datetime.now().astimezone().isoformat(),'pending_count':len(files),'ledgers':ledgers,'items':out,'warning':'States are reconciliation candidates, not permission to delete or rewrite ledgers.'}

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--settings',default=str(Path(__file__).parent.parent/'assets/settings.json'));p.add_argument('--mapping',required=True);p.add_argument('--history-dir',required=True);p.add_argument('--report',required=True)
    a=p.parse_args();settings=json.loads(Path(a.settings).read_text(encoding='utf-8-sig'));mapping=json.loads(Path(a.mapping).read_text(encoding='utf-8-sig'))
    output=Path(a.report).resolve()
    for key in ['pending_root','ledger_root','primary_archive_root','settlement_archive_root']:
        root=Path(settings[key]).resolve()
        if output==root or root in output.parents:raise ValueError('Write reconciliation report outside business roots')
    result=reconcile(settings,mapping,a.history_dir)
    output.write_text(json.dumps(result,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
    from collections import Counter
    print(json.dumps({'pending':result['pending_count'],'states':dict(Counter(i['status'] for i in result['items']))},ensure_ascii=False))
